import json
import os
import subprocess
import sys
import tempfile

def convert_to_docx(input_pdf, output_file):
    import docx
    from docx.shared import Inches, Pt

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    txt_result = subprocess.run(
        ['pdftotext', '-layout', input_pdf, '-'],
        capture_output=True,
        text=True,
        check=True
    )

    pages = txt_result.stdout.split('\x0c')
    for p_idx, page in enumerate(pages):
        lines = page.splitlines()
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.font.name = 'Arial'
            run.font.size = Pt(11)
        if p_idx < len(pages) - 1 and lines:
            doc.add_page_break()

    doc.save(output_file)

def convert_to_pptx(input_pdf, output_file):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory() as temp_dir:
        prefix = os.path.join(temp_dir, 'slide')
        subprocess.run(
            ['pdftoppm', '-png', '-r', '150', input_pdf, prefix],
            check=True
        )

        slide_files = sorted(
            [os.path.join(temp_dir, f) for f in os.listdir(temp_dir) if f.startswith('slide-') and f.endswith('.png')]
        )

        for img_path in slide_files:
            slide = prs.slides.add_slide(blank_slide_layout)
            slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=Inches(10), height=Inches(7.5))

    prs.save(output_file)

def convert_to_xlsx(input_pdf, output_file):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Page 1"

    txt_result = subprocess.run(
        ['pdftotext', '-layout', input_pdf, '-'],
        capture_output=True,
        text=True,
        check=True
    )

    pages = txt_result.stdout.split('\x0c')
    for p_idx, page in enumerate(pages):
        current_ws = ws if p_idx == 0 else wb.create_sheet(title=f"Page {p_idx + 1}")
        row_num = 1
        for line in page.splitlines():
            if not line.strip():
                continue
            cells = [c.strip() for c in line.split('   ') if c.strip()]
            if not cells:
                cells = [line.strip()]
            for col_num, val in enumerate(cells, start=1):
                current_ws.cell(row=row_num, column=col_num, value=val)
            row_num += 1

    wb.save(output_file)

def main():
    if len(sys.argv) < 4:
        print(json.dumps({'error': 'Uso: python converter_worker.py <format> <input> <output>', 'success': False}))
        sys.exit(1)

    target_format = sys.argv[1].lower()
    input_pdf = sys.argv[2]
    output_file = sys.argv[3]

    if not os.path.exists(input_pdf):
        print(json.dumps({'error': 'Archivo de entrada no encontrado', 'success': False}))
        sys.exit(1)

    try:
        if target_format == 'docx':
            convert_to_docx(input_pdf, output_file)
        elif target_format == 'pptx':
            convert_to_pptx(input_pdf, output_file)
        elif target_format == 'xlsx':
            convert_to_xlsx(input_pdf, output_file)
        else:
            print(json.dumps({'error': f'Formato no soportado: {target_format}', 'success': False}))
            sys.exit(1)

        print(json.dumps({'output': output_file, 'success': True}))
    except Exception as e:
        print(json.dumps({'error': str(e), 'success': False}))
        sys.exit(1)

if __name__ == '__main__':
    main()
